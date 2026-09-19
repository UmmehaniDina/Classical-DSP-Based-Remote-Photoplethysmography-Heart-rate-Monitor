from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT


REPORT = Path(r"C:\Users\dumme\Downloads\DSP Project\Section_6_Reflection_on_Individual_and_Team_Work.docx")


def put(cell, text):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.add_run(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


entries = [
    ("16 Jul 2026", "Collected, read, and discussed the main rPPG, CHROM, POS, detrending, and spectral-estimation papers needed to define the project direction.", "All members", "Joint literature preparation", "Established the DSP methods and limitations considered in the proposal."),
    ("20 Jul 2026", "Prepared the project proposal jointly and received proposal acceptance after confirming the software-only scope, input options, pipeline, and evaluation requirement.", "All members", "Proposal preparation and acceptance", "Formally initiated the approved classical DSP rPPG project."),
    ("22 Jul 2026", "Prepared the end-to-end processing pipeline from timestamped facial video through ROI extraction, RGB processing, candidate analysis, quality-gated BPM estimation, interface update, and saved evidence.", "All members", "Pipeline planning", "Established the serial module flow followed in the later project work."),
    ("23 Jul 2026", "Reviewed configuration structure for source selection, analysis-window duration, frequency limits, ROI settings, and output paths.", "Md. Montashim Billaha Chisty", "Configuration review", "Kept DSP assumptions explicit and reproducible."),
    ("24 Jul 2026", "Reviewed capture timestamps, monotonic fallback, duplicate handling, effective FPS, and unsafe-gap rejection.", "Mahir Nurain Shawchchow", "Timing review", "Protected interpolation and filtering from an invalid time base."),
    ("29 Jul 2026", "Reviewed face landmarks, forehead and cheek regions, coordinate smoothing, and normalized motion checks.", "Ummehani Dina", "ROI design review", "Connected stable spatial sampling to RGB-trace quality."),
    ("30 Jul 2026", "Reviewed YCrCb skin masking, minimum valid-pixel thresholds, and ROI validity metadata.", "Israt Ferdous Sabrin", "Extraction review", "Prevented weak or non-skin regions from entering the analysis buffer."),
    ("30 Jul 2026", "Reviewed percentile-clipped RGB pooling and fixed, adaptive-chroma, and temporal-adaptive masking alternatives.", "Md. Montashim Billaha Chisty", "RGB extraction review", "Reduced pixel outlier influence while preserving traceable extraction decisions."),
    ("31 Jul 2026", "Reviewed accepted-sample storage, frame events, coverage checks, and rolling-window requirements.", "All members", "Session-evidence review", "Separated acquisition failures from later low-signal rejections."),
    ("05 Aug 2026", "Reviewed uniform resampling, smoothness-priors detrending, FPS-aware lambda scaling, and channel normalization.", "Ummehani Dina", "Preprocessing review", "Prepared stable RGB signals for candidate generation."),
    ("06 Aug 2026", "Reviewed Green and PCA baselines and their candidate metadata.", "Israt Ferdous Sabrin", "Candidate-method review", "Established simple and data-driven reference candidates."),
    ("07 Aug 2026", "Reviewed local CHROM and POS projections with overlapping windows and local channel normalization.", "Md. Montashim Billaha Chisty", "Candidate-method review", "Added colour-projection methods to mitigate common illumination variation."),
    ("12 Aug 2026", "Reviewed band clamping and third-order Butterworth forward-reverse filtering.", "Mahir Nurain Shawchchow", "Filter design review", "Restricted analysis to the configured physiological frequency range without phase shift."),
    ("13 Aug 2026", "Reviewed Welch PSD settings, Hann tapering, overlap, spectral-peak refinement, and harmonic adjustment.", "Mahir Nurain Shawchchow", "Spectral-estimation review", "Defined the frequency-domain BPM estimate and safeguards."),
    ("14 Aug 2026", "Reviewed peak detection, polarity handling, median interval estimation, and robust outlier screening.", "Israt Ferdous Sabrin", "Time-domain estimation review", "Provided an independent BPM estimate for agreement checking."),
    ("19 Aug 2026", "Reviewed candidate selection, multi-feature quality scoring, and the role of coverage, motion, illumination, and spectral evidence.", "Israt Ferdous Sabrin", "Quality-control review", "Required both usable acquisition and a credible pulse signal before release."),
    ("20 Aug 2026", "Reviewed FFT-peak fusion, jump protection, EMA/Kalman tracking options, and low-quality rejection behaviour.", "Mahir Nurain Shawchchow", "Estimate-decision review", "Connected instantaneous estimates to a stable displayed BPM."),
    ("21 Aug 2026", "Reviewed session folders, metadata snapshots, RGB and BPM CSV records, frame events, and analysis events.", "All members", "Logging review", "Created the audit trail from input to accepted or rejected analysis."),
    ("26 Aug 2026", "Reviewed UBFC-style reference loading, overlap-only interpolation, paired comparisons, MAE, RMSE, and bias.", "All members", "Evaluation review", "Prepared a reproducible reference-comparison workflow."),
    ("27 Aug 2026", "Reviewed timing, extraction, DSP, tracking, session, evaluation, and GUI-oriented controlled tests.", "All members", "Regression review", "Checked the pipeline against controlled expected behaviour."),
    ("28 Aug 2026", "Reviewed command-line controls, headless replay, calibration mode, and the relationship between configuration and saved evidence.", "All members", "Application workflow review", "Allowed the same DSP core to be inspected outside the GUI."),
    ("02 Sep 2026", "Reviewed the background engine, bounded queues, GUI telemetry, and processing-to-interface data flow.", "All members", "Interactive application integration", "Kept the interactive view responsive while analysis operated on rolling windows."),
    ("03 Sep 2026", "Completed integration of the principal processing path: input, ROI extraction, preprocessing, candidate analysis, quality-gated BPM decision, session evidence, and interactive display.", "All members", "Core project completion", "The main DSP project and interactive workflow were substantially complete."),
    ("04 Sep 2026", "Reviewed warm-up, rejected, and accepted status presentation and the visibility of quality information in the application.", "All members", "Application improvement", "Improved the clarity of user-facing DSP state changes."),
    ("10 Sep 2026", "Inspected saved output sessions and compared unpaired webcam traces with the short paired UBFC-style replay.", "All members", "Result review", "Distinguished preliminary evaluation evidence from general accuracy claims."),
    ("11 Sep 2026", "Reviewed log completeness, application-side evidence flow, and small refinements to configuration and diagnostic behaviour.", "All members", "Application improvement", "Focused on minor integration and observability improvements after core completion."),
    ("14 Sep 2026", "Completed the final technical review of DSP settings, quality safeguards, interactive behaviour, session records, and limitations.", "All members", "Final technical review", "Last project-work date for the implemented software and DSP system."),
    ("18 Sep 2026", "Prepared final presentation slides and consolidated the report using the completed project evidence and development record.", "All members", "Submission preparation", "Presentation and report preparation completed after the technical work."),
]


doc = Document(REPORT)
table = doc.tables[0]
while len(table.rows) > 1:
    table._tbl.remove(table.rows[-1]._tr)
for entry in entries:
    cells = table.add_row().cells
    for cell, value in zip(cells, entry):
        put(cell, value)
doc.save(REPORT)
print(REPORT)
