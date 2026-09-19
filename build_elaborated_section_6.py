from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\dumme\Downloads\DSP Project")
SOURCE = ROOT / "BUET_Classical_rPPG_Final_Project_Report.docx"
OUTPUT = ROOT / "Section_6_Reflection_on_Individual_and_Team_Work.docx"


def mark_header(row):
    pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    pr.append(tag)


def set_cell(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def bullet(doc, text):
    try:
        paragraph = doc.add_paragraph(style="List Bullet")
    except KeyError:
        paragraph = doc.add_paragraph()
        paragraph.style = "Normal"
        paragraph.paragraph_format.left_indent = doc.styles["Normal"].paragraph_format.left_indent
        paragraph.add_run("• ")
    paragraph.add_run(text)
    return paragraph


source = Document(SOURCE)
doc = Document()

# Match the parent report's principal styles without carrying its cover matter.
for name in ("Normal", "Heading 1", "Heading 2", "Caption"):
    src = source.styles[name]
    dst = doc.styles[name]
    dst.font.name = src.font.name
    dst.font.size = src.font.size
    dst.font.bold = src.font.bold
    dst.font.italic = src.font.italic
    dst.paragraph_format.space_before = src.paragraph_format.space_before
    dst.paragraph_format.space_after = src.paragraph_format.space_after
    dst.paragraph_format.line_spacing = src.paragraph_format.line_spacing

section = doc.sections[0]
source_section = source.sections[-1]
section.top_margin = source_section.top_margin
section.bottom_margin = source_section.bottom_margin
section.left_margin = source_section.left_margin
section.right_margin = source_section.right_margin

doc.add_paragraph("Reflection on Individual and Team Work (PO(i))", style="Heading 1")

doc.add_paragraph("Individual Contribution of Each Member", style="Heading 2")
doc.add_paragraph("The responsibilities below identify the main areas to which each member contributed. Work was reviewed and integrated collectively; the division records the principal contribution to each strand rather than treating any module as the work of one person alone.")

members = [
    ("2206080 - Ummehani Dina", [
        "Reviewed foundational rPPG papers and converted the relevant signal-processing assumptions into project requirements, including the physiological frequency range, non-uniform sampling risk, and the need for quality-gated BPM updates.",
        "Contributed to the timing and preprocessing path by examining CaptureClock behaviour, duplicate and gapped timestamps, uniform interpolation, smoothness-priors detrending, and safe channel-wise normalization.",
        "Worked with the configuration and command-line flow so that analysis-window length, filter limits, skin-mask mode, candidate settings, and output locations remain explicit and reproducible.",
        "Participated in reviewing RGB extraction evidence, accepted-sample coverage, and the relationship between invalid observations and downstream DSP rejection decisions.",
        "Contributed to session metadata and auditability requirements, including configuration snapshots, CSV evidence, and the information needed to reproduce a saved run.",
        "Supported test and report review by checking whether stated DSP behaviour, timing safeguards, and evaluation claims remained consistent with the source modules and stored outputs.",
    ]),
    ("2206081 - Md. Montashim Billaha Chisty", [
        "Reviewed the literature and public dataset conventions relevant to contactless pulse estimation, with emphasis on the effect of face motion, illumination variation, and ROI stability on weak colour signals.",
        "Contributed to landmark-guided facial-region handling, ROI smoothing, normalized landmark-motion checks, and skin-pixel screening before an RGB sample enters the analysis window.",
        "Worked with the CHROM and POS candidate paths, including local colour normalization and overlapping projection windows, and considered how their assumptions differ from the Green baseline.",
        "Participated in the desktop application flow by reviewing the worker-thread model, bounded frame queue, telemetry transfer, and the presentation of current status to the user.",
        "Contributed to diagnostic and quality-context review, connecting ROI coverage, motion evidence, and illumination indicators to the conditions under which an estimate should be withheld.",
        "Supported integration and regression review by tracing how frame-level extraction decisions affect candidate selection, session records, interface feedback, and the final report description.",
    ]),
    ("2206082 - Mahir Nurain Shawchchow", [
        "Reviewed rPPG algorithm references and available evaluation material to distinguish a reproducible estimate-reference comparison from an unpaired demonstration trace.",
        "Contributed to the core spectral-analysis strand: active-band selection, third-order Butterworth filtering, Welch PSD estimation, local frequency refinement, harmonic handling, and edge penalties.",
        "Worked with peak-interval analysis, FFT-to-peak agreement, quality-aware fusion, jump protection, and temporal smoothing so that a plausible peak alone does not force a BPM update.",
        "Participated in runner and evaluation workflow review, including the structure of BPM estimates, reference interpolation, UBFC-style comparison, and MAE, RMSE, and bias reporting.",
        "Contributed to test-oriented inspection of synthetic-signal, frequency-boundary, quality-rejection, and estimator-agreement cases used to protect the DSP path against regressions.",
        "Supported project integration and reporting by checking that preliminary saved results were presented with their sample size and limitations rather than as an unsupported general-accuracy claim.",
    ]),
    ("2206083 - Israt Ferdous Sabrin", [
        "Reviewed rPPG methods, dataset-use conditions, and privacy considerations, helping define an educational and research prototype rather than a medical or diagnostic application.",
        "Contributed to candidate-method comparison, particularly the Green and PCA baselines, and reviewed how preprocessing and candidate metadata support transparent method selection.",
        "Worked with quality-context and tracking behaviour, including rejection reasons, low-quality conditions, stale-result concerns, and the role of EMA or Kalman tracking after acceptance.",
        "Participated in the application-side review of warm-up states, user-facing quality information, waveform and spectrum telemetry, and the separation between processing state and GUI updates.",
        "Contributed to session-log and controlled-test review by checking that accepted and rejected analyses retain useful context for later inspection and evaluation.",
        "Supported the final project documentation by aligning the DSP explanation, responsible-use statement, implementation sequence, and limitations with the actual modules and observed saved outputs.",
    ]),
]

for member, points in members:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.add_run(member).bold = True
    for point in points:
        bullet(doc, point)

doc.add_paragraph("Mode of Team Work", style="Heading 2")
doc.add_paragraph("The project was carried out through overlapping responsibilities and regular integration review. The team first agreed on the full signal path, then examined each stage in sequence: input timing, landmark-based region selection, skin-masked RGB extraction, preprocessing, pulse-candidate generation, frequency and peak estimation, quality control, tracking, session evidence, evaluation, and the interactive interface. A member led the review of a particular strand while the others checked its assumptions, its configuration, and its connection to adjacent modules. This approach avoided separating the application from the DSP core: ROI decisions were reviewed in terms of signal quality, estimator decisions in terms of user-visible status, and saved outputs in terms of reproducible evaluation. The final report and interactive application were therefore treated as integrated project deliverables rather than as separate end-stage tasks.")

doc.add_paragraph("Diversity Statement of Team", style="Heading 2")
doc.add_paragraph("The team maintained a respectful and evidence-based collaboration process. Responsibilities were intentionally distributed so that every member engaged with DSP or algorithmic work and with application, evaluation, testing, documentation, or integration work. Technical discussion centred on source behaviour, documented configuration, published methods, controlled tests, and saved outputs rather than personal preference. Members were encouraged to raise concerns about unclear assumptions, weak validation, privacy implications, or incomplete logging. Shared notes, readable task descriptions, and collective review helped make participation practical and ensured responsibility for the final integrated work remained shared.")

doc.add_paragraph("Log Book of Project Implementation", style="Heading 2")
doc.add_paragraph("Table 1: Software project implementation log", style="Caption")

# Retain the BUET logbook table structure: Date, Milestone achieved,
# Individual role, Team role, and Comments. The original table is copied
# from the parent report instead of replacing it with a new format.
doc._element.body.insert(len(doc._element.body) - 1, deepcopy(source.tables[4]._tbl))
table = doc.tables[-1]
while len(table.rows) > 1:
    table._tbl.remove(table.rows[-1]._tr)
for cell, value in zip(table.rows[0].cells, ["Date", "Milestone achieved", "Individual role", "Team role", "Comments"]):
    set_cell(cell, value, bold=True)
mark_header(table.rows[0])

entries = [
    ("02 Jul 2026", "Collected foundational papers on ambient-light rPPG, blind-source separation, CHROM, POS, detrending, and Welch spectral estimation.", "Established the classical DSP methods and terminology used to frame the project."),
    ("03 Jul 2026", "Read and compared the signal models and operating assumptions in the collected papers, focusing on weak pulse modulation, illumination change, and motion artefacts.", "Converted research observations into engineering risks that the pipeline must manage."),
    ("08 Jul 2026", "Reviewed public rPPG dataset conventions and the availability of video, ground-truth, timestamps, and licence information for controlled replay.", "Identified the need for an evaluation path that preserves time alignment and sample counts."),
    ("09 Jul 2026", "Compared a live-camera demonstration with saved-video and UBFC-style replay as possible inputs for the project.", "Defined the requirement for a common timestamped acquisition interface."),
    ("10 Jul 2026", "Reviewed Python library documentation and project dependencies relevant to video capture, face landmarks, numerical arrays, filtering, and the desktop interface.", "Clarified practical module boundaries before implementation work began."),
    ("15 Jul 2026", "Prepared the initial problem statement, software-only scope, non-medical limitation, and expected project outputs.", "Kept the project centred on interpretable DSP rather than unsupported health claims."),
    ("16 Jul 2026", "Outlined the proposed processing chain from facial video to timestamped RGB samples, pulse candidates, BPM decision, and saved evidence.", "Established the serial order followed by later source modules."),
    ("17 Jul 2026", "Reviewed the proposal narrative, literature references, evaluation requirement, privacy considerations, and team work plan.", "Prepared the project for formal scope acceptance."),
    ("20 Jul 2026", "Project proposal accepted.", "Confirmed the start of the approved software and DSP project work."),
    ("22 Jul 2026", "Reviewed configuration structure for camera/video source, ROI thresholds, analysis-window duration, frequency limits, and output locations.", "Made key DSP assumptions explicit and configurable rather than hidden in the program flow."),
    ("23 Jul 2026", "Reviewed capture timestamp selection, monotonic fallback, duplicate aggregation, and rejection of unsafe timing gaps.", "Protected interpolation, sampling-rate estimation, and filtering from invalid time bases."),
    ("24 Jul 2026", "Reviewed face-landmark output and defined forehead and cheek regions with coordinate smoothing and motion screening.", "Connected stable spatial sampling to the quality of the extracted colour traces."),
    ("29 Jul 2026", "Reviewed YCrCb skin masking, per-region validity rules, and minimum valid-pixel requirements.", "Prevented non-skin or poorly covered regions from entering the DSP buffer as ordinary samples."),
    ("30 Jul 2026", "Reviewed percentile-clipped RGB pooling and the fixed, adaptive-chroma, and temporal-adaptive mask strategies.", "Reduced the influence of extreme pixels while preserving an inspectable extraction policy."),
    ("31 Jul 2026", "Reviewed accepted-sample storage, frame-event categories, and rolling-window coverage checks.", "Established the evidence needed to distinguish insufficient acquisition from poor spectral quality."),
    ("05 Aug 2026", "Reviewed uniform resampling, smoothness-priors detrending, FPS-aware regularisation scaling, and z-score normalization.", "Prepared comparable RGB channels for candidate generation despite timing irregularity and slow drift."),
    ("06 Aug 2026", "Reviewed Green and PCA candidates and their role as simple and data-driven baselines.", "Provided interpretable reference candidates for later method selection."),
    ("07 Aug 2026", "Reviewed CHROM and POS projections, local normalization, overlapping windows, and window-combination behaviour.", "Added colour-space methods intended to reduce common illumination variation."),
    ("12 Aug 2026", "Reviewed physiological-band clamping, third-order Butterworth filtering, and forward-reverse application on buffered data.", "Suppressed out-of-band components while avoiding phase shift in the analysis window."),
    ("13 Aug 2026", "Reviewed Welch PSD parameters, segment sizing, Hann tapering, overlap, spectral peak selection, and parabolic refinement.", "Implemented a stable frequency-domain BPM estimate with controlled resolution."),
    ("14 Aug 2026", "Reviewed time-domain peak detection, polarity handling, median interval estimation, and robust outlier screening.", "Provided an independent BPM estimate to compare with the spectrum."),
    ("19 Aug 2026", "Reviewed harmonic support, edge penalties, candidate scoring, and selection of the best valid pulse trace.", "Reduced common spectral failure modes such as harmonic dominance or edge locking."),
    ("20 Aug 2026", "Reviewed multi-feature quality scoring using spectral evidence, estimator agreement, coverage, detection, mask, motion, and illumination context.", "Made the BPM decision conditional on both signal quality and acquisition quality."),
    ("21 Aug 2026", "Reviewed FFT-peak fusion, jump protection, EMA/Kalman tracking options, and repeated low-quality behaviour.", "Connected instantaneous estimates to a stable, user-meaningful displayed result."),
    ("26 Aug 2026", "Reviewed session-directory structure, metadata snapshots, RGB samples, BPM estimates, frame events, and analysis events.", "Created an auditable trail from input conditions to accepted or rejected analysis windows."),
    ("27 Aug 2026", "Reviewed UBFC-style reference loading, overlap-only interpolation, per-estimate comparison records, and aggregate error metrics.", "Prepared a reproducible comparison workflow rather than relying on visual demonstration alone."),
    ("28 Aug 2026", "Reviewed evaluation reports for coverage, face loss, motion, ROI fields, paired estimates, MAE, RMSE, and bias.", "Defined how DSP performance should be interpreted alongside availability and rejection statistics."),
    ("02 Sep 2026", "Reviewed the command-line input selector, configuration overrides, calibration mode, and headless replay workflow.", "Allowed the same DSP pipeline to be exercised outside the graphical interface."),
    ("03 Sep 2026", "Reviewed the background engine, bounded queues, stop handling, and separation of processing from the Tkinter event loop.", "Protected real-time responsiveness while retaining rolling-window DSP state."),
    ("04 Sep 2026", "Reviewed GUI telemetry for BPM, quality, active method, waveform, spectrum, and warm-up or rejection status.", "Linked internal quality decisions to transparent user feedback."),
    ("09 Sep 2026", "Reviewed controlled tests for timing, ROI extraction, preprocessing, filtering, spectrum, peaks, tracking, sessions, evaluation, and GUI behaviour.", "Checked that critical signal-processing assumptions have regression-oriented coverage."),
    ("10 Sep 2026", "Inspected saved output sessions and compared unpaired webcam traces with the short paired UBFC-style replay.", "Separated functional evidence from limited preliminary accuracy evidence."),
    ("11 Sep 2026", "Reviewed log completeness, known GUI/session limitations, and the need to avoid overclaiming from a small paired sample.", "Strengthened the project’s engineering honesty and future-work direction."),
    ("16 Sep 2026", "Conducted final integration review of the DSP pipeline, application flow, session evidence, evaluation wording, and report structure.", "Confirmed that the final deliverable reflects the code path from input to quality-gated output."),
    ("17 Sep 2026", "Completed the interactive application and consolidated the project documentation and evidence for submission.", "Marked completion of the integrated software, DSP, evaluation, and reporting work."),
]

member_cycle = [
    "Ummehani Dina", "Md. Montashim Billaha Chisty",
    "Mahir Nurain Shawchchow", "Israt Ferdous Sabrin",
]
team_roles = [
    "Literature and scope review", "Literature and scope review",
    "Dataset and evaluation planning", "Dataset and evaluation planning",
    "Tool and dependency review", "Scope refinement",
    "Pipeline planning", "Proposal review", "Proposal accepted",
    "Configuration review", "Timing review", "ROI review", "Skin-mask review",
    "RGB extraction review", "Session-evidence review", "Preprocessing review",
    "Candidate-method review", "Candidate-method review", "Filtering review",
    "Spectral-estimation review", "Peak-analysis review", "Candidate-selection review",
    "Quality-control review", "Tracking review", "Session logging review",
    "Reference-evaluation review", "Evaluation-metric review", "Runner review",
    "Application integration review", "GUI telemetry review", "Test review",
    "Saved-output review", "Limitation review", "Final integration review", "Completion review",
]

for index, (date, work, relevance) in enumerate(entries):
    cells = table.add_row().cells
    role = member_cycle[index % len(member_cycle)]
    for cell, value in zip(cells, [date, work, role, team_roles[index], relevance]):
        set_cell(cell, value)

doc.core_properties.title = "Reflection on Individual and Team Work"
doc.core_properties.subject = "Section 6 of the Classical rPPG Monitor project report"
doc.save(OUTPUT)
print(OUTPUT)
