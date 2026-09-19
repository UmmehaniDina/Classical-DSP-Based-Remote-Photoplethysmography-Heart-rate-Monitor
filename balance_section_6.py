from pathlib import Path

from docx import Document


REPORT = Path(r"C:\Users\dumme\Downloads\DSP Project\Section_6_Reflection_on_Individual_and_Team_Work.docx")


def replace_paragraph(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


doc = Document(REPORT)

# Align individual contribution bullets to the final named technical logbook entries.
for paragraph in doc.paragraphs:
    if paragraph.text.startswith("Contributed to the timing and preprocessing path"):
        replace_paragraph(paragraph, "Contributed to the landmark-guided ROI and preprocessing strand, including forehead and cheek region review, uniform resampling requirements, smoothness-priors detrending, and safe normalization of the RGB traces.")
    elif paragraph.text.startswith("Worked with the configuration and command-line flow"):
        replace_paragraph(paragraph, "Worked with command-line configuration flow, session metadata, and output handling so that analysis-window length, filter limits, candidate settings, and output locations remain explicit and reproducible.")
    elif paragraph.text.startswith("Contributed to landmark-guided facial-region handling"):
        replace_paragraph(paragraph, "Contributed to configuration and extraction setup, including source selection, analysis settings, RGB pooling, skin-pixel screening, and the checks applied before a sample enters the analysis window.")
    elif paragraph.text.startswith("Contributed to the core spectral-analysis strand"):
        replace_paragraph(paragraph, "Contributed to the timing and spectral-analysis strand: progressing timestamps, effective sampling rate, active-band selection, third-order Butterworth filtering, Welch PSD estimation, local frequency refinement, harmonic handling, and edge penalties.")

# Balance named logbook responsibilities: four named technical entries per member.
table = doc.tables[0]
for row in table.rows[1:]:
    date = row.cells[0].text
    milestone = row.cells[1].text
    if date == "31 Jul 2026" and milestone.startswith("Reviewed accepted-sample storage"):
        row.cells[2].text = "Ummehani Dina"
        row.cells[3].text = "Session-evidence implementation"
    elif date == "28 Aug 2026" and milestone.startswith("Reviewed command-line controls"):
        row.cells[2].text = "Ummehani Dina"
        row.cells[3].text = "Command-line workflow implementation"
    elif date == "02 Sep 2026" and milestone.startswith("Reviewed the background engine"):
        row.cells[2].text = "Md. Montashim Billaha Chisty"
        row.cells[3].text = "Interactive application implementation"

doc.save(REPORT)
print(REPORT)
