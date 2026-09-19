from pathlib import Path
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\dumme\Downloads\DSP Project")
REF = Path(r"C:\Users\dumme\Downloads\EEE-xxx-project-report-template (1).docx")
OUT = ROOT / "BUET_Classical_rPPG_Final_Project_Report.docx"
ASSETS = ROOT / "report_work" / "report_figures"
ASSETS.mkdir(parents=True, exist_ok=True)


FONT = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 27)
SMALL = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 22)
TITLE = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 28)

def centered(draw, box, text, font, fill="#0f172a"):
    lines = text.split("\n")
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + 5 * (len(lines)-1)
    y = (box[1]+box[3]-total)/2
    for line, height in zip(lines, heights):
        width = draw.textbbox((0,0), line, font=font)[2]
        draw.text(((box[0]+box[2]-width)/2, y), line, font=font, fill=fill)
        y += height+5

def architecture_figure(path):
    img = Image.new("RGB", (2100, 900), "white"); d = ImageDraw.Draw(img)
    labels = ["Video input\nand timestamps", "Face landmarks\nand multi ROI", "Skin-masked\nRGB traces", "DSP candidates\nand quality gates", "BPM estimate,\nGUI, session logs"]
    colors = ["#dbeafe", "#dcfce7", "#fef3c7", "#fce7f3", "#e0e7ff"]
    for i, (label, color) in enumerate(zip(labels, colors)):
        x = 45 + i*410; box = (x, 280, x+300, 535)
        d.rounded_rectangle(box, radius=25, fill=color, outline="#334155", width=4)
        centered(d, box, label, FONT)
        if i < 4:
            d.line((x+305, 408, x+390, 408), fill="#334155", width=5)
            d.polygon([(x+390,408),(x+370,394),(x+370,422)], fill="#334155")
    title = "Rolling-window, quality-aware software pipeline"
    w=d.textbbox((0,0),title,font=TITLE)[2]; d.text(((2100-w)/2, 690), title, font=TITLE, fill="#0f172a")
    img.save(path)


def dsp_figure(path):
    img = Image.new("RGB", (1800, 1100), "white"); d = ImageDraw.Draw(img)
    d.rectangle((150, 100, 1690, 480), outline="#475569", width=3)
    d.rectangle((150, 620, 1690, 1000), outline="#475569", width=3)
    d.text((150, 40), "Illustrative preprocessing trace", font=TITLE, fill="#0f172a")
    d.text((150, 560), "Illustrative spectral decision", font=TITLE, fill="#0f172a")
    xs=np.linspace(150,1690,650)
    raw=290+105*np.sin(np.linspace(0, 4*np.pi, 650))+50*np.sin(np.linspace(0, 28*np.pi,650))
    clean=290+45*np.sin(np.linspace(0, 28*np.pi,650))
    d.line(list(zip(xs,raw)), fill="#64748b", width=3)
    d.line(list(zip(xs,clean)), fill="#2563eb", width=4)
    d.text((1150,130), "Raw trace", font=SMALL, fill="#64748b"); d.text((1150,165), "Detrended trace", font=SMALL, fill="#2563eb")
    low=150+1540*55/240; high=150+1540*200/240
    d.rectangle((low,620,high,1000), fill="#dcfce7")
    d.line((150,960,1690,960), fill="#334155", width=3)
    points=[]
    for x in np.linspace(150,1690,650):
        bpm=(x-150)*240/1540
        amp=35+260*np.exp(-((bpm-72)/11)**2)+50*np.exp(-((bpm-144)/20)**2)
        points.append((x,960-amp))
    d.line(points, fill="#7c3aed", width=4)
    peakx=150+1540*72/240; d.line((peakx,640,peakx,960), fill="#dc2626", width=3)
    d.text((peakx+15,655), "72 BPM", font=SMALL, fill="#dc2626")
    d.text((low+15,930), "Accepted search band", font=SMALL, fill="#166534")
    img.save(path)


def evaluation_figure(path):
    img=Image.new("RGB",(2100,800),"white"); d=ImageDraw.Draw(img)
    stages=["Permitted\nvideo / dataset","Fixed\nconfiguration","Session CSV\noutputs","Reference BPM\nalignment","MAE, RMSE, bias,\nacceptance metrics"]
    for i,stage in enumerate(stages):
        x=45+i*410; box=(x,220,x+300,460)
        d.rounded_rectangle(box, radius=25, fill="#f8fafc", outline="#0f172a", width=3); centered(d,box,stage,FONT)
        if i<4:
            d.line((x+305,340,x+390,340),fill="#334155",width=5); d.polygon([(x+390,340),(x+370,326),(x+370,354)],fill="#334155")
    text="Reproducible evaluation requires matched clocks, documented settings, and review of rejected windows."
    w=d.textbbox((0,0),text,font=SMALL)[2]; d.text(((2100-w)/2,650),text,font=SMALL,fill="#0f172a")
    img.save(path)


architecture_figure(ASSETS / "architecture.png")
dsp_figure(ASSETS / "dsp.png")
evaluation_figure(ASSETS / "evaluation.png")


def set_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]: r.text = ""
    else:
        p.add_run(text)


def p(doc, text, style="Normal", page=False):
    par = doc.add_paragraph(style=style)
    if page: par.add_run().add_break(WD_BREAK.PAGE)
    par.add_run(text)
    return par


def h(doc, text, level, page=False): return p(doc, text, f"Heading {level}", page)
def cap(doc, text): return p(doc, text, "Caption")


def fig(doc, image, caption_text, width=6.1):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = par.add_run().add_picture(str(image), width=Inches(width))
    shape._inline.docPr.set("descr", caption_text)
    cap(doc, caption_text)


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for cell, v in zip(t.rows[0].cells, headers): cell.text = v
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row in rows:
        cells = t.add_row().cells
        for cell, v in zip(cells, row): cell.text = v
    return t


def toc(sdt):
    content = sdt.find(qn("w:sdtContent"))
    for child in list(content): content.remove(child)
    entries = [
        "Table of Contents", "Abstract", "Introduction", "Design", "    Problem Formulation (PO(b))",
        "    Design Method (PO(a))", "    System Architecture", "    Algorithm Design", "    Software Implementation",
        "    Real-Time Execution and Profiling", "Implementation", "Design Analysis and Evaluation",
        "Reflection on Individual and Team Work (PO(i))", "Communication to External Stakeholders (PO(j))",
        "Project Management (PO(k))", "Future Work (PO(l))", "References"
    ]
    for i, entry in enumerate(entries):
        para = OxmlElement("w:p"); props = OxmlElement("w:pPr")
        if i == 0:
            style = OxmlElement("w:pStyle"); style.set(qn("w:val"), "TOCHeading"); props.append(style)
        para.append(props); run = OxmlElement("w:r"); tx = OxmlElement("w:t"); tx.text = entry; run.append(tx); para.append(run); content.append(para)


doc = Document(REF)
set_text(doc.paragraphs[3], "[Course Code] ([Term])")
set_text(doc.paragraphs[4], "[Course Title]")
set_text(doc.paragraphs[6], "Section: [Section]   Group: [Group]")
set_text(doc.paragraphs[8], "Classical rPPG Monitor for Contactless Heart Rate Estimation")
set_text(doc.paragraphs[9], "Prepared by: [Student Name 1] ([Student ID 1])")
set_text(doc.paragraphs[10], "             [Student Name 2] ([Student ID 2])")
set_text(doc.paragraphs[14], "Course Instructor / Supervisor:")
set_text(doc.paragraphs[15], "[Supervisor Name], [Designation]")
set_text(doc.paragraphs[16], "")
set_text(doc.paragraphs[21], "In signing this statement, we certify that this report and the software work it describes are our own. Literature, algorithms, datasets, and software libraries used in the project are cited. Student signatures: [Student 1] ____________   [Student 2] ____________")
body = doc._element.body
children = list(body)
toc(children[25])
for child in children[27:-1]: body.remove(child)

# ABSTRACT
h(doc, "Abstract", 1, True)
p(doc, "This report documents the design and implementation of a software-only classical remote photoplethysmography (rPPG) monitor developed for the final project. The application estimates heart rate from facial video by extracting a weak pulse-related colour variation and processing it with an explicitly configured DSP chain. A live camera stream, saved video, or UBFC-style recording is timestamped; facial landmarks define forehead and cheek regions; valid skin pixels are pooled into RGB samples; and a rolling window is resampled before detrending, colour projection, band-pass filtering, spectral estimation, peak analysis, and estimate fusion. The implemented candidates are Green, PCA, CHROM, and POS. A result is released only when the timing, sample coverage, spatial evidence, motion, illumination, spectral prominence, and agreement checks are adequate. The program also includes a Tkinter desktop interface, a command-line runner, session CSV records, and reference-alignment utilities. The report is deliberately limited to claims supported by the source code and saved project outputs. The available UBFC-style replay contains only three paired estimates and is consequently reported as preliminary evidence rather than as a general accuracy result. The system is an educational and research prototype, not a medical device; it must not be used for diagnosis, treatment, or emergency decisions.")

# INTRODUCTION
h(doc, "Introduction", 1)
p(doc, "Heart rate is commonly measured by contact sensors, but a video-based estimate can be useful when a non-contact educational demonstration or retrospective video analysis is preferred. The target signal is extremely weak: normal facial movement, ambient-light changes, automatic image adjustments, compression, and missed landmarks can be much larger than the pulse-related variation. A credible solution therefore requires more than a single frequency peak. It must create a trace from meaningful facial regions, preserve timing information, evaluate competing pulse models, and reject observations that lack sufficient evidence.")
p(doc, "The project adopts a classical DSP approach. Unlike a trained end-to-end model, this approach has visible intermediate signals and a direct relationship between the engineering assumptions and the outcome. The green channel is a useful baseline because hemoglobin absorption makes its relative variation informative under suitable light [1]. CHROM and POS form color-space projections that seek to reduce common-mode illumination variation [3], [4]. PCA provides an additional data-driven candidate. Each method can fail in different conditions, so the system compares candidates rather than treating one method as universally reliable.")
p(doc, "The problem is a complex engineering problem because the input is non-stationary, the sample timing is imperfect, the desired component is small, and the outcome must remain understandable when it is rejected. The constraints include a rolling analysis window, physiological search limits, absence of reliable input during face loss, and a user interface that should never imply false certainty. The project therefore emphasizes traceability: accepted RGB samples, rejected frame events, quality features, and BPM estimates are written to a session directory for review.")
p(doc, "The report is framed as a final course project. It retains the course-report sections required by the supplied template while focusing exclusively on software design, signal processing, evaluation, project management, and responsible use. It does not include thesis-only front matter or any physical design material.")

# DESIGN
h(doc, "Design", 1)
h(doc, "Problem Formulation (PO(b))", 2)
h(doc, "Identification of Scope", 3)
p(doc, "The software accepts a permitted live video stream, a saved video file, or a compatible recorded dataset folder. It outputs an estimated BPM only after an accepted analysis window. The system is responsible for landmark-guided region selection, skin-pixel screening, RGB extraction, temporal preprocessing, candidate construction, filtering, spectral and peak estimation, quality control, smoothing, display, and session evidence. It is not intended to measure oxygen saturation, diagnose disease, infer emotion, identify people, or make a medical decision.")
h(doc, "Literature Review", 3)
p(doc, "Verkruysse, Svaasand, and Nelson showed that pulse-related variation can be observed in ordinary ambient-light video [1]. Poh, McDuff, and Picard introduced blind-source separation for non-contact cardiac pulse measurement [2]. De Haan and Jeanne proposed CHROM, which uses chrominance information to improve robustness [3]. Wang et al. described the POS algorithmic principle [4]. Tarvainen et al. introduced smoothness-priors detrending for physiological time series [5], while Welch described averaged periodograms for more stable power-spectrum estimates [6]. These works motivate the selectable, inspectable components of the present project.")
h(doc, "Formulation of Problem", 3)
p(doc, "Let c[k] = [R[k], G[k], B[k]]^T denote a timestamped RGB observation extracted from accepted skin pixels. The desired scalar pulse component p[k] is obscured by illumination I[k], non-pulsatile reflectance, movement, and noise. A simplified observation model is c[k] = I[k](u_s s[k] + u_d d_0) + u_p p[k] + n[k], where the terms represent common reflection, static diffuse color, pulse-related color variation, and noise. The software seeks a BPM value b = 60f from a dominant in-band frequency f while requiring temporal, spatial, and spectral evidence sufficient to accept the result.")
h(doc, "Analysis", 3)
p(doc, "The analysis window is not assumed to be uniformly sampled. Timestamps are sorted, duplicate times are aggregated, excessive gaps are rejected, and finite RGB observations are linearly interpolated onto a uniform grid. If the median interval is delta t, the effective sampling rate is f_s = 1/delta t. The selected heart-rate band is 55 to 200 BPM, or approximately 0.917 to 3.333 Hz. The upper edge is clamped below Nyquist so that a low-rate stream cannot silently produce an invalid filter or spectrum.")
p(doc, "The core decision is intentionally conservative. A valid spectrum alone is insufficient if the recent window has poor sample coverage or unstable illumination. Likewise, a peak-interval value is not accepted unconditionally. The final quality score contains a spectral term, FFT-to-peak agreement, accepted-sample coverage, face-detection availability, skin-mask coverage, motion quality, and illumination quality. The result is rejected if a required timing or quality condition fails.")

h(doc, "Design Method (PO(a))", 2)
p(doc, "The design applies mathematics, signal processing, and software engineering in a connected sequence. Temporal sampling theory determines whether a received stream can represent the accepted pulse band. Robust statistics reduce the influence of bright or dark outlier pixels. Regularized detrending separates slow baseline variation from the target pulse range. Colour projection creates several candidate pulse traces. Frequency and time-domain estimates provide complementary evidence. Finally, an event-driven program architecture preserves the evidence needed to explain why a window was accepted, warming up, or rejected. This makes the system suitable for technical inspection: an estimate can be traced back to its source timestamps, accepted RGB samples, selected candidate, quality features, and configuration snapshot.")
p(doc, "For each color channel y, smoothness-priors detrending estimates a slowly varying trend z by solving z_hat = (I + lambda^2 D_2^T D_2)^(-1)y, where D_2 is a second-difference matrix and lambda controls trend smoothness [5]. The residual y - z_hat is standardized before candidate generation. The implementation uses a sparse factorization cache for repeated window lengths, which avoids constructing the same linear solve on every update.")
p(doc, "After candidate construction, a third-order Butterworth band-pass filter is applied in forward and reverse directions. The effective magnitude is squared while phase distortion is cancelled in the buffered analysis window. Welch PSD analysis averages overlapping tapered segments, which reduces variance compared with a single raw periodogram [6]. Local parabolic interpolation around the strongest spectral bin refines the frequency estimate without reporting a value outside the configured band.")

h(doc, "System Architecture", 2)
p(doc, "The system is organized into five software layers. The acquisition layer obtains frames and selects a progressing timestamp source. The perception layer returns face landmarks and derives named regions. The extraction layer applies ROI masks and produces one robust RGB observation only when sufficient valid pixels are present. The DSP layer maintains the rolling sample window, builds candidate signals, calculates quality measures, and emits an accepted or rejected analysis result. The presentation and persistence layer updates the user-facing telemetry and writes metadata and CSV records.")
cap(doc, "Table 1: Software layers and their responsibilities")
table(doc, ["Layer", "Main responsibility", "Evidence produced"], [
    ["Acquisition", "Read frames and normalize timestamps", "Frame timing and source label"],
    ["Perception", "Locate landmarks and derive facial regions", "Region geometry and motion measure"],
    ["Extraction", "Select valid skin pixels and compute RGB means", "RGB sample and mask coverage"],
    ["DSP", "Create, filter, score, and fuse pulse candidates", "BPM candidates and quality features"],
    ["Presentation", "Update GUI and write session artifacts", "Status, plots, CSV, metadata"],
])

h(doc, "Algorithm Design", 2)
h(doc, "Region Selection and RGB Extraction", 3)
p(doc, "The face is represented by landmarks from which the forehead, left cheek, and right cheek regions are derived. The regions are smoothed over time to reduce coordinate jitter. A normalized landmark displacement is calculated between consecutive frames. When displacement exceeds the configured motion threshold, the RGB observation is excluded from the pulse window. This prevents a large facial movement from being treated as a color oscillation.")
p(doc, "Within each region, pixels are converted to YCrCb. A fixed chrominance mask uses configurable Y, Cr, and Cb limits. The software also provides current-frame adaptive chroma and temporal adaptive variants, but the mask strategy remains explicit in the session record. Per-channel robust means use values clipped between configured low and high percentiles. The extraction stage returns non-finite RGB when the total valid-pixel requirement is not met; it does not substitute arbitrary fallback color values.")
p(doc, "The default configuration is deliberately specific. It requires at least five valid skin pixels per region and twenty in total. Its fixed YCrCb gate uses Y > 35, 133 <= Cr <= 180, and 77 <= Cb <= 135. The RGB mean is calculated after 10th-to-90th-percentile clipping. Region coordinates are exponentially smoothed with alpha = 0.45, and normalized landmark movement above 0.025 rejects the corresponding sample. These values are not claimed to be universal physiological constants; they are explicit engineering settings that can be reviewed, changed, and preserved in session metadata.")
h(doc, "Candidate Pulse Signals", 3)
p(doc, "The Green candidate is the standardized green trace. PCA decomposes the normalized three-channel matrix and evaluates each component by spectral quality and explained variance. Its sign can be aligned against the green candidate to make inspection more consistent. CHROM forms X = 3R - 2G and Y = 1.5R + G - 1.5B, then combines them using an adaptive standard-deviation weight. POS forms S1 = G - B and S2 = -2R + G + B with the corresponding adaptive combination. CHROM and POS are evaluated in overlapping local windows after local channel normalization so that their coefficients respond to recent conditions rather than a single global mean.")
h(doc, "Filtering, Spectrum, and Peak Estimates", 3)
p(doc, "Every valid candidate is filtered in the active search band. The system begins with a global band or forms an adaptive band around the last reliable BPM. A periodic global-band probe is retained so that the adaptive search does not remain locked to a stale local range. Welch analysis returns frequency bins and power values. The peak frequency is refined, clipped to the active band, checked for a plausible subharmonic, and penalized when it lies too close to a band edge. Independent peak detection measures the median inter-peak period in the filtered waveform and optionally evaluates both polarities.")
p(doc, "The default analysis window is 24 s with a 16 s warm-up requirement. The accepted physiological interval is 55--200 BPM. The third-order Butterworth filter is applied using forward--reverse filtering, while the Welch segment is configured from an 8 s duration, a minimum of 64 samples, a 0.10 Hz target resolution, a Hann taper, and overlapping segments. Peak detection uses a BPM-derived minimum distance and robust median-absolute-deviation screening. These settings balance spectral resolution, response time, and numerical safety for the declared operating range.")
h(doc, "Estimate Fusion and Temporal Tracking", 3)
p(doc, "The spectral BPM is the anchor estimate. If the peak BPM is finite, it receives a bounded configurable share that decreases continuously as its disagreement with the spectral BPM grows. A jump guard rejects sudden changes relative to the previous accepted BPM. The runner then applies either a time-aware exponential moving average or an optional BPM Kalman tracker. The design requirement is to clear stale data after repeated low-quality windows; inspection of the GUI path shows that this policy needs explicit completion there before a production release.")
p(doc, "The implemented quality score is not a single arbitrary threshold. Its spectral term is multiplied by an agreement factor between FFT and peak estimates and by an acquisition term derived from sample coverage, face-detection availability, mask coverage, motion quality, and illumination quality. The default minimum spectral ratio is 2.0. The fusion step limits the peak contribution to 0.30 and uses an exponential agreement scale of 18 BPM; the default jump guard is 18 BPM. Consequently, a numerically plausible peak is not sufficient when the input window is incomplete, the face is unstable, or the two estimators disagree.")

h(doc, "Software Implementation", 2)
p(doc, "The implementation uses a Python source-layout package. Configuration is represented by validated dataclasses that group camera/video input settings, ROI settings, and analysis settings. The DSP package contains preprocessing, filtering, spectral estimation, and peak utilities. The algorithms package contains candidate generation, analysis, and temporal tracking. Separate modules manage ROI geometry, RGB extraction, landmarks, timing, session writing, calibration evidence, evaluation, and visualization.")
p(doc, "In the source tree, preprocessing.py implements sparse smoothness-priors detrending and caches the sparse factorization by window length and regularisation value. It scales the detrending parameter with the square of the effective sampling-rate ratio, preventing the same nominal setting from changing meaning when the camera rate changes. spectral.py implements the band clamp, Welch spectrum, local frequency refinement, harmonic adjustment, and edge penalty. analysis.py is responsible for candidate selection, multi-feature quality computation, and fusion. This separation is important because it localises DSP assumptions instead of burying them in the user-interface code.")
p(doc, "The command-line runner maintains deques for accepted samples and all frame events. The accepted-sample deque is pruned to the configured analysis-window duration; the frame-event deque supports coverage and detection checks over the same window. Each analysis event has an incrementing sequence number, a status, a rejection category when applicable, sample count, timing metrics, and accepted-result metadata. This structure supports replay, debugging, and reproducible inspection without coupling the analysis logic to the GUI. The desktop engine shares the analysis core, but its session writer should be extended to record frame and analysis-event rows with the same completeness as the runner.")
p(doc, "The desktop interface is separated from the analysis engine through a background worker and queues. Frame delivery favors recency by dropping an outdated display frame when the queue is full. Telemetry messages contain only the values needed for the current UI update, including BPM, quality, selected method, waveform, spectrum, and status. This separation keeps the event loop responsive while the DSP code operates on the latest completed window.")

h(doc, "Real-Time Execution and Profiling", 2)
p(doc, "Real-time behavior is managed through bounded work rather than an assumption that every machine has the same speed. Face processing occurs per frame, while full-window analysis is scheduled at a configurable interval. The window is normally 24 seconds and the update interval is one second. A warm-up requirement prevents early analysis. If accepted samples become stale, the program clears the displayed BPM and resets its tracker. For offline video, the GUI engine paces playback using the declared frame rate where it is credible so that visual replay remains understandable.")
p(doc, "The GUI uses a worker thread and bounded queues. The frame queue has capacity two and favours a recent frame when the consumer cannot keep pace; this is appropriate for display because latency is more harmful than losing an obsolete image. Analysis state remains in the engine rather than being shared with the Tkinter event loop. The design therefore separates time-critical acquisition and DSP work from periodic user-interface updates. Measurements of speed are intentionally not reported because the repository does not contain a controlled profiling run for the current computer, input resolution, and dependency set.")
p(doc, "Profiling should record separate costs for landmark detection, ROI extraction, preprocessing, candidate generation, filtering, spectral estimation, peak detection, and UI dispatch. The repository provides timing and session metrics but a final performance figure depends on the actual execution environment, input resolution, source codec, and configured methods. The report therefore gives a reproducible profiling method rather than presenting unverified latency as a universal claim.")

# IMPLEMENTATION
h(doc, "Implementation", 1)
h(doc, "Description", 2)
p(doc, "A session begins by validating configuration and creating an output directory with logs, metadata, CSV files, and figure locations. The application opens the configured input, initializes the landmark model, and begins collecting frame events. During normal operation it shows a warm-up state until the data window contains enough recent accepted RGB samples. After a valid analysis update, the status contains BPM and the latest quality measure. When quality is low, the interface keeps the reason visible rather than silently continuing the previous display.")
h(doc, "Threading and State Management", 2)
p(doc, "The background engine owns frame acquisition and DSP state. The GUI polls queues for the newest image and events, avoiding direct mutation of analysis buffers from the display thread. Stop requests are represented by an event checked inside the processing loop. Resources are closed in a finalization block so that a video capture, landmarker, dashboard, and CSV writers are released after normal completion or failure. The same design is used for command-line processing, with the presentation portion omitted in headless mode.")
h(doc, "Configuration and Reproducibility", 2)
p(doc, "The default TOML configuration documents requested input dimensions, target frame rate, ROI thresholds, analysis window, filter order, accepted BPM limits, peak settings, spectral settings, smoothing behavior, and quality limits. The command line exposes selected overrides, such as input source, ROI strategy, analysis-window duration, candidate method, and output directory. Each session stores a configuration snapshot, allowing a later reader to reproduce the stated setting rather than infer it from the code version.")
h(doc, "Verification Through Automated Tests", 2)
p(doc, "The repository contains 83 named automated test functions spanning configuration validation, timestamp behavior, ROI clamping, RGB extraction, temporal masks, spectral refinement, filtering, fusion, candidate selection, tracking, session CSV schemas, UBFC reference parsing, evaluation, and GUI components. These tests express behavior on controlled inputs: for example, a synthetic sinusoid should be estimated near its known BPM, a timestamp gap should cause rejection, and an invalid skin mask should not create an RGB sample. The current project virtual environment points to an unavailable interpreter, so a fresh full-suite pass is not claimed in this report. The inventory is evidence of intended regression coverage, not a substitute for reference-device validation on real video.")
cap(doc, "Table 2: Verification matrix for the software pipeline")
table(doc, ["Verification target", "Controlled input", "Expected behavior"], [
    ["Timestamp handling", "Duplicate, backward, or gapped times", "Aggregate duplicates; reject unsafe gaps"],
    ["Signal processing", "Synthetic known-frequency trace", "Estimate remains within configured tolerance"],
    ["ROI extraction", "Invalid or low-coverage region", "Reject RGB sample without fallback"],
    ["Candidate selection", "Finite and non-finite candidate traces", "Prefer accepted high-quality candidate"],
    ["Session evidence", "Accepted and rejected analyses", "Write stable CSV fields and categories"],
    ["User interface", "Engine lifecycle and widget state", "Remain responsive and report state"],
])

# ANALYSIS AND EVALUATION
h(doc, "Design Analysis and Evaluation", 1)
h(doc, "Novelty", 2)
p(doc, "The project does not claim a novel physiological principle. Its contribution is a disciplined integration of established classical methods into one inspectable software workflow. The integration includes landmark-aware multi-region extraction, robust pixel pooling, timing-aware interpolation, multiple color projections, spectral and time-domain agreement, periodic global-band recovery, quality-context logging, and a user interface that distinguishes a valid estimate from a poor-quality window. This is valuable in a course setting because every major decision can be traced to a configurable rule or recorded session field.")
h(doc, "Design Considerations (PO(c))", 2)
h(doc, "Considerations to Public Health and Safety", 3)
p(doc, "A numerical BPM display can be misinterpreted as a clinical reading. The application therefore describes its output as an estimate, uses rejection states, and does not offer diagnostic recommendations. The software must not be used for triage, medication decisions, emergency assessment, or any setting where a qualified measurement is required. Quality gating reduces obvious misuse of poor input; it does not establish medical-grade accuracy.")
h(doc, "Considerations to Environment", 3)
p(doc, "The project is software-only and does not require project-specific consumables. Its environmental footprint is mainly ordinary compute and display use. Bounded buffers, replay-based debugging, and configurable logging can reduce unnecessary repeated processing and storage. These benefits are modest and should not be overstated.")
h(doc, "Considerations to Cultural and Societal Needs", 3)
p(doc, "Facial-video methods can perform differently with lighting, camera processing, motion, facial hair, occlusion, compression, and appearance. The design avoids claiming equal performance for all people or conditions. It records coverage and rejection signals so that a failed or weak observation can be recognized. Evaluation should use authorized data, report condition-specific limitations, and avoid any deployment that turns video observation into unconsented monitoring.")
h(doc, "Investigations (PO(d))", 2)
h(doc, "Design of Experiment", 3)
p(doc, "A controlled evaluation replays a permitted recorded sequence with fixed configuration and compares each accepted BPM with a timestamp-aligned reference series. The protocol first records frame count, effective FPS, detection availability, sample acceptance, and rejection categories. It then reports the number of valid estimate-reference pairs. Only after these availability measures are known should mean absolute error (MAE), root mean squared error (RMSE), mean bias, and correlation be interpreted. A simple definition is MAE = (1/N) sum |b_est,i - b_ref,i| and RMSE = sqrt((1/N) sum (b_est,i - b_ref,i)^2).")
h(doc, "Data Collection", 3)
p(doc, "The software supports an UBFC-rPPG folder layout containing a video and a ground-truth text file. It reads finite BPM and timestamps, normalizes timestamp units where needed, and writes comparison records only over the overlapping time interval. The saved session metadata identifies the UBFC-style input used at the time of execution. No new human-subject data are claimed in this report. Any new recording should be performed with informed consent, an approved collection procedure where applicable, and storage controls appropriate to facial video.")
h(doc, "Results and Analysis", 3)
p(doc, "The inspected output directory contains two kinds of preliminary evidence. A webcam session retained twelve accepted BPM updates between 109.335 and 110.267 BPM; it has no paired reference series and is therefore a functional trace, not an accuracy result. A short UBFC-style replay retained three timestamp-aligned estimate-reference pairs. Its saved summary reports mean estimate 109.335 BPM against mean reference 113.333 BPM, MAE 3.998 BPM, RMSE 4.026 BPM, and bias -3.998 BPM. These values demonstrate that the comparison path executed, but N = 3 is far too small to support a general performance claim. A larger controlled evaluation requires a repaired environment, declared input identities, configuration snapshots, and acceptance statistics.")
cap(doc, "Table 3: Required reporting fields for a reproducible evaluation")
table(doc, ["Metric", "Meaning", "Report requirement"], [
    ["Paired estimates", "Number of in-range estimated/reference pairs", "State count before error metrics"],
    ["MAE", "Average absolute BPM deviation", "State unit BPM and window policy"],
    ["RMSE", "Error magnitude with larger-error emphasis", "State unit BPM and pairing method"],
    ["Bias", "Mean signed BPM deviation", "Identify positive or negative direction"],
    ["Coverage", "Accepted samples relative to received frames", "Report with face-detection availability"],
    ["Rejections", "Why analysis windows were rejected", "Report categories and counts"],
])
p(doc, "Table 3A makes the small sample size of the saved replay visible. These values were read from its stored summary and must be replaced with a larger controlled evaluation before any general accuracy conclusion is made.")
cap(doc, "Table 3A: Preliminary saved UBFC-style replay evidence (not a general accuracy claim)")
table(doc, ["Quantity", "Saved value", "Interpretation"], [
    ["Paired estimate count", "3", "Insufficient for broad conclusion"],
    ["Mean estimate", "109.335 BPM", "Accepted values in saved replay"],
    ["Mean reference", "113.333 BPM", "Time-aligned reference values"],
    ["MAE / RMSE", "3.998 / 4.026 BPM", "Preliminary error magnitudes only"],
    ["Mean bias", "-3.998 BPM", "Underestimate in this short replay"],
])
h(doc, "Interpretation and Conclusions on Data", 3)
p(doc, "An error metric without context can be misleading. A low MAE from a small, stable subset of a video does not prove robustness during face loss or motion. Conversely, a high rejection rate may protect against false BPM updates but lower the fraction of time for which the system is useful. The appropriate interpretation combines availability, rejection reasons, and paired-estimate accuracy. This report therefore treats the evaluation harness as a core deliverable and reserves strong performance claims for a completed, reproducible experiment.")
h(doc, "Limitations of Tools (PO(e))", 2)
p(doc, "The software depends on input-timestamp quality, landmark reliability, library versions, and operating-system scheduling. Automatic exposure or white balance behavior may produce intensity changes that a classical projection cannot fully remove. Low illumination raises noise and longer exposure can blur motion. Major pose changes or occlusions can prevent stable ROI extraction. A fixed 30 FPS target is adequate for basic heart-rate estimation in the accepted band but does not support fine waveform morphology or precise pulse-transit analysis. The saved environment’s missing interpreter target is a reproducibility limitation. Inspection also found two completion items: GUI-originated sessions should write frame and analysis-event records, and a unique session identifier should prevent output-directory collision when sessions begin in the same second.")
h(doc, "Impact Assessment (PO(f))", 2)
h(doc, "Assessment of Societal and Cultural Issues", 3)
p(doc, "A contactless estimate can be attractive in education, wellness demonstrations, and non-contact interaction, but a face stream is sensitive information. Users should know when capture is active, whether frames are stored, and where session evidence is written. The application should never be repurposed for covert observation or evaluation of a person without consent.")
h(doc, "Assessment of Health and Safety Issues", 3)
p(doc, "The primary safety risk is over-trust in a value that is inaccurate or stale. The mitigation is software behavior and user communication: state a non-medical purpose, reset stale data, reject low-quality windows, expose status, and avoid clinical language. Users must seek qualified care for health concerns rather than relying on this application.")
h(doc, "Assessment of Legal Issues", 3)
p(doc, "Use of face video, reference data, and public datasets must respect consent, license terms, institutional requirements, and applicable privacy law. Session directories may contain timestamps and derived physiological estimates, so access should be limited and retention should be deliberate. The report references established algorithms and identifies third-party libraries and datasets rather than presenting them as original work.")
h(doc, "Sustainability Evaluation (PO(g))", 2)
p(doc, "The system’s sustainability consideration is primarily efficient software operation. It uses a bounded rolling window, a scheduled analysis interval, and session-local outputs. Such choices limit unnecessary memory growth and repeated computation. Any broader sustainability claim would require measurement of deployment-specific energy use, which is outside the present project scope.")
h(doc, "Ethical Issues (PO(h))", 2)
p(doc, "Ethical design requires that a rejected condition remain visible instead of being hidden behind a plausible-looking number. The project preserves that distinction through quality gates and logs. It also requires consent for recording, careful treatment of facial-video data, and honesty about unequal performance across conditions. The project intentionally avoids claims of clinical validation and does not use individual identity as a feature of the analysis.")

# TEAM / STAKEHOLDERS / MANAGEMENT
h(doc, "Reflection on Individual and Team Work (PO(i))", 1)
h(doc, "Individual Contribution of Each Member", 2)
p(doc, "[Student Name 1] - [Replace with actual contribution: literature review, DSP design, implementation, testing, or documentation].")
p(doc, "[Student Name 2] - [Replace with actual contribution: GUI, evaluation, session logging, testing, or documentation].")
h(doc, "Mode of Team Work", 2)
p(doc, "The team should divide work by module boundaries, agree on configuration and data contracts before integration, and review changes against automated tests and documented expected behavior. A practical workflow is to assign one member responsibility for a module while keeping shared ownership of interfaces, validation, report writing, and final demonstration. Replace this statement with the team’s actual working process before submission.")
h(doc, "Diversity Statement of Team", 2)
p(doc, "The team commits to respectful collaboration, equitable participation, accessible communication, and evidence-based technical discussion. Work allocation and review should be based on declared responsibilities and available time, not personal characteristics. Design decisions should be challenged through requirements, data, and tests.")
h(doc, "Log Book of Project Implementation", 2)
cap(doc, "Table 4: Software project implementation log")
table(doc, ["Date", "Milestone achieved", "Individual role", "Team role", "Comments"], [
    ["[Date]", "Scope, requirements, and literature reviewed", "[Name]", "Planning", "Replace with actual record"],
    ["[Date]", "ROI and RGB extraction integrated", "[Name]", "Implementation", "Replace with actual record"],
    ["[Date]", "DSP candidate and quality pipeline completed", "[Name]", "Integration", "Replace with actual record"],
    ["[Date]", "GUI, session logging, and tests reviewed", "[Name]", "Validation", "Replace with actual record"],
])

h(doc, "Communication to External Stakeholders (PO(j))", 1)
h(doc, "Executive Summary", 2)
p(doc, "This project is a desktop software tool that estimates a pulse-related frequency from facial video using established signal-processing methods. It tracks facial regions, converts valid skin-pixel colors into a time signal, compares several classical rPPG methods, and withholds updates when the data are unreliable. The program saves evidence for later review. It is intended for learning and research demonstration, not medical care. Any wider use would require consented evaluation, clear privacy controls, and independent validation against an appropriate reference measurement.")
h(doc, "User Manual", 2)
p(doc, "1. Prepare a valid Python environment and install the project dependencies. 2. Start the desktop application or command-line runner. 3. Choose a permitted live source, saved video, or compatible dataset folder. 4. Keep the face visible in steady illumination and minimize motion during warm-up. 5. Wait for the analysis window to fill. 6. Treat a BPM value as usable only when the status indicates an accepted quality result. 7. Stop the session normally. 8. Review the created session directory for configuration metadata, frame events, RGB samples, BPM estimates, and analysis events. 9. Do not use the output for diagnosis or urgent decisions.")
h(doc, "Source Repository Link", 2)
p(doc, "[Insert authorized source repository URL]")
h(doc, "Demonstration Video Link", 2)
p(doc, "[Insert authorized demonstration video URL]")

h(doc, "Project Management (PO(k))", 1)
h(doc, "Timeline of Project Implementation", 2)
p(doc, "The project is managed as a sequence of software deliverables: requirements and literature review, architecture definition, ROI and extraction development, DSP implementation, GUI/session integration, automated testing, controlled replay evaluation, and report review. Each phase should have an owner, a review point, and a clear artifact. A risk register should include unavailable dependencies, missing input permissions, unstable video timing, and unverified benchmark claims.")
cap(doc, "Table 5: Software project timeline and deliverables")
table(doc, ["Phase", "Activity", "Deliverable", "Primary risk"], [
    ["1", "Requirements and literature review", "Approved scope and citations", "Overstated health claim"],
    ["2", "Architecture and configuration", "Module contract and defaults", "Unclear data flow"],
    ["3", "ROI, extraction, and DSP integration", "Runnable core pipeline", "Weak input quality"],
    ["4", "GUI, logging, and tests", "Auditable session workflow", "Thread/state mismatch"],
    ["5", "Replay evaluation and reporting", "Reproducible evidence", "Missing environment or reference alignment"],
])

h(doc, "Future Work (PO(l))", 1)
p(doc, "The immediate next step is to repair and lock the Python environment, run the automated tests, and record exact dependency versions. Controlled replay should then be performed on authorized data with known reference timestamps. Future software work can improve GUI logging parity, add a profiling report, provide a configuration wizard, display quality factors more clearly, and make privacy retention configurable. Algorithmic research could investigate stronger motion compensation, confidence calibration, broader condition-specific evaluation, and carefully validated learned models. Heart-rate variability, respiration estimation, and oxygen-related measures are separate research topics and must not be represented as features of the present application without independent validation.")

h(doc, "References", 1)
refs = [
    "[1] W. Verkruysse, L. O. Svaasand, and J. S. Nelson, \"Remote plethysmographic imaging using ambient light,\" Optics Express, vol. 16, no. 26, pp. 21434-21445, 2008.",
    "[2] M.-Z. Poh, D. J. McDuff, and R. W. Picard, \"Non-contact, automated cardiac pulse measurements using video imaging and blind source separation,\" Optics Express, vol. 18, no. 10, pp. 10762-10774, 2010.",
    "[3] G. de Haan and V. Jeanne, \"Robust pulse rate from chrominance-based rPPG,\" IEEE Transactions on Biomedical Engineering, vol. 60, no. 10, pp. 2878-2886, 2013.",
    "[4] W. Wang, A. C. den Brinker, S. Stuijk, and G. de Haan, \"Algorithmic principles of remote PPG,\" IEEE Transactions on Biomedical Engineering, vol. 64, no. 7, pp. 1479-1491, 2017.",
    "[5] M. P. Tarvainen, P. O. Ranta-aho, and P. A. Karjalainen, \"An advanced detrending method with application to HRV analysis,\" IEEE Transactions on Biomedical Engineering, vol. 49, no. 2, pp. 172-175, 2002.",
    "[6] P. Welch, \"The use of fast Fourier transform for the estimation of power spectra,\" IEEE Transactions on Audio and Electroacoustics, vol. 15, no. 2, pp. 70-73, 1967.",
    "[7] C. Lugaresi et al., \"MediaPipe: A framework for building perception pipelines,\" arXiv:1906.08172, 2019.",
    "[8] S. Bobbia, R. Macwan, Y. Benezeth, A. Mansouri, and J. Dubois, \"Unsupervised skin tissue segmentation for remote photoplethysmography,\" Pattern Recognition Letters, vol. 124, pp. 82-90, 2019.",
    "[9] SciPy Community, \"SciPy Signal Processing documentation,\" https://docs.scipy.org/doc/scipy/reference/signal.html, accessed September 2026.",
    "[10] OpenCV, \"OpenCV documentation,\" https://docs.opencv.org/, accessed September 2026.",
]
for r in refs: p(doc, r)

# Ask Word to refresh future TOC/PAGE fields.
settings = doc.settings.element
u = settings.find(qn("w:updateFields"))
if u is None:
    u = OxmlElement("w:updateFields"); settings.append(u)
u.set(qn("w:val"), "true")
# Add non-visual descriptions without changing the template's decorative drawing geometry.
for drawing in doc._element.xpath(".//wp:docPr"):
    if not drawing.get("descr"):
        drawing.set("descr", "Template decorative element")
doc.save(OUT)
print(OUT)
