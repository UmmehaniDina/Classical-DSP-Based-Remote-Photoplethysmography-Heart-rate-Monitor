from pathlib import Path

from docx import Document


path = Path(r"C:\Users\dumme\Downloads\EEE-xxx-project-report-template (1).docx")
doc = Document(path)
for i, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.replace("\t", "[TAB]")
    if text.strip() or (paragraph._p.pPr is not None and paragraph._p.pPr.sectPr is not None):
        style = paragraph.style.name
        fmt = paragraph.paragraph_format
        has_sect = paragraph._p.pPr is not None and paragraph._p.pPr.sectPr is not None
        print(f"P {i:03d} | {style} | sect={has_sect} | align={paragraph.alignment} | before={fmt.space_before} after={fmt.space_after} | {text}")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
    for ri, row in enumerate(table.rows):
        print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
