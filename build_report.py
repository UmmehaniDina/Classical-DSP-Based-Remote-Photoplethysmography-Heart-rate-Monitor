from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\dumme\Downloads\DSP Project")
REFERENCE = Path(r"C:\Users\dumme\Downloads\EEE-xxx-project-report-template (1).docx")
OUTPUT = ROOT / "Classical_rPPG_Monitor_Final_Project_Report.docx"


def set_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def add_para(doc, text="", style="Normal", page_break=False):
    p = doc.add_paragraph(style=style)
    if page_break:
        p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run(text)
    return p


def add_heading(doc, text, level, page_break=False):
    return add_para(doc, text, f"Heading {level}", page_break)


def add_caption(doc, text):
    return add_para(doc, text, "Caption")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = width
    return table


def build_toc_sdt(sdt):
    content = sdt.find(qn("w:sdtContent"))
    for child in list(content):
        content.remove(child)

    entries = [
        "Table of Contents",
        "Abstract",
        "Introduction",
        "Design",
        "    Problem Formulation (PO(b))",
        "    Design Method (PO(a))",
        "    System Architecture",
        "    Algorithm Design",
        "    Software Implementation",
        "    Real-Time Execution and Profiling",
        "Implementation",
        "Design Analysis and Evaluation",
        "Reflection on Individual and Team Work (PO(i))",
        "Communication to External Stakeholders (PO(j))",
        "Project Management (PO(k))",
        "Future Work (PO(l))",
        "References",
    ]
    for i, entry in enumerate(entries):
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        if i == 0:
            pstyle = OxmlElement("w:pStyle")
            pstyle.set(qn("w:val"), "TOCHeading")
            ppr.append(pstyle)
        p.append(ppr)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = entry
        r.append(t)
        p.append(r)
        content.append(p)


doc = Document(REFERENCE)

# Cover-page placeholders. Existing paragraph/run formatting remains in place.
set_text(doc.paragraphs[3], "[Course Code] ([Term])")
set_text(doc.paragraphs[4], "[Course Title]")
set_text(doc.paragraphs[6], "Section: [Section]   Group: [Group]")
set_text(doc.paragraphs[8], "Classical rPPG Monitor for Contactless Heart-Rate Estimation")
set_text(doc.paragraphs[9], "Prepared by: [Student Name 1] ([Student ID 1])")
set_text(doc.paragraphs[10], "             [Student Name 2] ([Student ID 2])")
set_text(doc.paragraphs[14], "Course Instructor / Supervisor:")
set_text(doc.paragraphs[15], "[Supervisor Name], [Designation]")
set_text(doc.paragraphs[16], "")
set_text(doc.paragraphs[18], "Signature of Instructor:   ___________________________________________________")
set_text(doc.paragraphs[21], "In signing this statement, we certify that this software project report is our own work. Sources, algorithms, and software libraries used in the project are acknowledged in the references. Student signatures: [Student 1] ____________   [Student 2] ____________")

# Preserve cover, section boundaries, and the TOC container. Remove all template body content after the TOC.
body = doc._element.body
children = list(body)
toc = children[25]
build_toc_sdt(toc)
# Body child 26 is the second explicit section-boundary paragraph. Retain it
# so the reference's three-section page system and first-page behavior stay intact.
for child in children[27:-1]:
    body.remove(child)

# Report body.
add_heading(doc, "Abstract", 1, page_break=True)
add_para(doc, "This project develops a software-only classical remote photoplethysmography (rPPG) monitor for contactless heart-rate estimation from a camera stream or saved video. The system uses face landmarks to define forehead and cheek regions, extracts robust RGB traces from skin pixels, resamples irregular observations, and evaluates Green, PCA, CHROM, and POS pulse candidates. It applies smoothness-priors detrending, Butterworth band-pass filtering, Welch spectral analysis, peak-interval estimation, quality gating, and time-aware BPM tracking. The application includes a desktop interface, a headless replay path, session CSV logging, UBFC-rPPG support, and unit tests for configuration, signal processing, timing, extraction, tracking, evaluation, and GUI behavior. The project is a research and educational prototype; its output is not medical advice or a diagnostic measurement.")

add_heading(doc, "Introduction", 1)
add_para(doc, "Heart-rate estimation is useful for observing periodic physiological variation, but conventional sensors require physical contact. Remote photoplethysmography estimates a weak blood-volume-related color variation in facial video under ambient illumination [1]. In practice, the desired signal is small relative to motion, illumination changes, compression artifacts, and missed face detections. The engineering problem is therefore not only selecting a pulse frequency; it is constructing a dependable software pipeline that rejects unreliable observations and reports its uncertainty.")
add_para(doc, "The project addresses this problem with an interpretable classical DSP design instead of a trained end-to-end model. Alternatives include a green-channel baseline, chrominance projections, blind-source separation, and learned video models. The selected design keeps intermediate signals, quality measures, and rejection causes available for inspection. This is appropriate for a course project because the processing choices can be tested independently and the limitations are explicit.")

add_heading(doc, "Design", 1)
add_heading(doc, "Problem Formulation (PO(b))", 2)
add_heading(doc, "Identification of Scope", 3)
add_para(doc, "Input is a live camera stream or an offline video sequence. The software detects facial landmarks, samples color from landmark-guided regions, estimates a heart-rate candidate in a configured physiological band, and stores accepted and rejected analysis events. The scope excludes clinical validation, diagnosis, and any claim of medical-grade accuracy.")
add_heading(doc, "Literature Review", 3)
add_para(doc, "Ambient-light rPPG was demonstrated through spatially averaged skin color changes in video [1]. CHROM uses chrominance combinations to reduce common illumination variation [2], while POS applies a plane-orthogonal-to-skin projection [3]. Smoothness-priors detrending is used to suppress slow baseline variation before pulse-band analysis [4]. These methods motivate the project’s multi-candidate, quality-gated approach.")
add_heading(doc, "Formulation of Problem", 3)
add_para(doc, "Given timestamped RGB observations x(t) from accepted facial skin pixels, estimate a BPM value b such that b = 60f, where f is the dominant, quality-supported frequency in the accepted search band. The estimate is accepted only when timing, coverage, illumination stability, spectral prominence, candidate agreement, and plausible BPM transitions meet configured criteria.")
add_heading(doc, "Analysis", 3)
add_para(doc, "Raw frame intervals are irregular and some frames are rejected. The system first interpolates valid RGB observations onto a uniform time base. Each channel is detrended and standardized. Candidate pulse signals are filtered in the target band and scored by a Welch spectrum. Peak intervals provide an independent BPM estimate. A fused estimate is accepted only after quality and jump guards; otherwise the system records a low-quality outcome rather than fabricating a value.")

add_heading(doc, "Design Method (PO(a))", 2)
add_para(doc, "The design applies signal-processing, statistics, and software engineering together. A median frame interval defines effective sampling frequency. The global search band is derived from the configured BPM limits and clamped below Nyquist. A zero-phase Butterworth filter removes slow trend residue and high-frequency noise. Welch estimation identifies spectral power in the band, while local parabolic interpolation refines the dominant bin. The quality score combines spectral prominence, FFT-to-peak agreement, sample coverage, face-detection availability, skin-mask coverage, motion quality, and illumination quality.")

add_heading(doc, "System Architecture", 2)
add_para(doc, "The architecture has five software layers: acquisition, landmark and ROI processing, RGB extraction, DSP analysis, and presentation/session storage. The acquisition layer supplies frame timestamps. Landmark processing produces named facial regions and motion evidence. Extraction produces RGB samples only when skin-pixel thresholds are met. The DSP layer processes a sliding analysis window. The presentation layer exposes status, waveform and spectrum data, and records CSV and metadata files for later evaluation.")
add_table(doc, ["Layer", "Responsibility", "Output"], [
    ["Acquisition", "Read live or replayed frames and normalize timestamps", "Frame and monotonic time"],
    ["ROI processing", "Detect landmarks, smooth regions, measure motion", "Validated facial regions"],
    ["Extraction", "Mask skin pixels and calculate robust RGB means", "Timestamped RGB sample"],
    ["DSP analysis", "Resample, filter, score candidates, fuse BPM", "Accepted BPM or rejection"],
    ["Presentation", "Display telemetry and write session evidence", "GUI status and CSV records"],
])

add_heading(doc, "Algorithm Design", 2)
add_para(doc, "Four classical candidates are evaluated: the green channel, principal components of normalized RGB, CHROM, and POS. CHROM and POS use locally normalized overlapping windows so that their color projections remain responsive to short-term changes. The implementation optionally enables ICA when its dependency is available, but it is not required for normal operation. Candidate selection chooses the accepted result with the strongest quality score while periodically probing the global band to reduce local-band lock-in.")
add_caption(doc, "Table 1: Principal DSP stages and their purpose")
add_table(doc, ["Stage", "Method", "Purpose"], [
    ["Timing", "Timestamp validation and interpolation", "Create an evenly sampled analysis signal"],
    ["Preprocessing", "Smoothness-priors detrending and z-score", "Reduce baseline drift and scale differences"],
    ["Filtering", "Zero-phase Butterworth band-pass", "Retain plausible pulse-band content"],
    ["Frequency estimation", "Welch spectrum and peak refinement", "Estimate dominant pulse frequency"],
    ["Cross-check", "Peak intervals and estimate fusion", "Reduce single-method dependence"],
    ["Quality control", "Coverage, motion, illumination and SNR gates", "Reject unreliable windows"],
])

add_heading(doc, "Software Implementation", 2)
add_para(doc, "The project is implemented in Python with a source-layout package. Separate modules manage configuration, ROI geometry, extraction, timing, filtering, spectral analysis, candidate generation, tracking, evaluation, session writing, and the GUI. Configuration is represented by typed dataclasses and validated before acquisition. The implementation stores candidate metadata and quality features alongside BPM estimates, making each accepted output auditable during replay or debugging.")
add_para(doc, "The desktop UI runs processing in a background engine and communicates telemetry through queues. The worker captures frames, updates the rolling window, and emits only the latest display frame when the consumer is slower. This prevents the interface from blocking the acquisition loop while keeping visual feedback current.")

add_heading(doc, "Real-Time Execution and Profiling", 2)
add_para(doc, "The software uses a bounded sliding window rather than retaining a full session in memory. Analysis is scheduled at a configurable interval, while frame-level extraction continues continuously. The project records effective FPS, face-detection availability, sample coverage, rejection reasons, and analysis outcomes. These measures identify whether a degraded result originates from cadence, visibility, motion, illumination, or weak pulse evidence. Performance claims are intentionally limited: runtime depends on the camera, video codec, processor, MediaPipe runtime, and display load.")

add_heading(doc, "Implementation", 1)
add_heading(doc, "Software Pipeline and Threading", 2)
add_para(doc, "For each frame, the system obtains a progressing timestamp, detects facial landmarks, computes named regions, and extracts a robust RGB observation. If the frame passes ROI and motion checks, it enters the analysis deque. At each update interval, the application validates frame-window quality, performs DSP analysis, and updates the BPM tracker only when the fused estimate is finite and accepted. Rejected windows preserve their reason in the session log, allowing the user to distinguish warm-up from poor signal conditions.")
add_heading(doc, "User Interface and Session Logging", 2)
add_para(doc, "The GUI presents acquisition status, current BPM, quality information, waveform and spectrum views, and operational telemetry. The session writer creates metadata plus CSV files for accepted RGB observations, frame events, BPM estimates, and analysis events. Offline UBFC-rPPG folders can be replayed and compared with their available reference series. The user may choose a camera or saved video source through the interface or use the command-line runner for headless processing.")

add_heading(doc, "Design Analysis and Evaluation", 1)
add_heading(doc, "Novelty", 2)
add_para(doc, "The project’s contribution is an auditable integration of classical rPPG methods rather than a claim of a new physiological algorithm. It combines multi-candidate estimation, landmark-aware skin sampling, explicit timing checks, adaptive-band recovery, quality-context logging, and a desktop presentation layer in one software-only teaching prototype.")
add_heading(doc, "Design Considerations (PO(c))", 2)
add_heading(doc, "Considerations to Public Health and Safety", 3)
add_para(doc, "The interface identifies the output as an estimated heart rate from a research prototype. It must not be used for diagnosis, triage, or treatment. The quality gates reduce the chance of presenting a number during obvious signal failure, but they do not establish clinical accuracy.")
add_heading(doc, "Considerations to Environment", 3)
add_para(doc, "The software has no project-specific consumables. Its environmental impact is limited mainly to ordinary computation and display energy use. Offline replay and bounded data retention help avoid unnecessary repeated acquisition and storage.")
add_heading(doc, "Considerations to Cultural and Societal Needs", 3)
add_para(doc, "Facial-video processing can perform differently across lighting conditions, cameras, facial movement, skin appearance, and video compression. The design records mask coverage and rejected samples instead of assuming uniform performance. Evaluation should use consented data and report condition-specific limitations honestly.")
add_heading(doc, "Investigations (PO(d))", 2)
add_heading(doc, "Design of Experiment", 3)
add_para(doc, "A reproducible evaluation consists of replaying a permitted recorded video or UBFC-rPPG subject folder with fixed configuration, collecting the generated CSV files, and comparing accepted estimates with the aligned reference BPM series where available. The same session should record acceptance rate, rejection reasons, and paired-estimate count before interpreting error values.")
add_heading(doc, "Data Collection", 3)
add_para(doc, "The repository includes code for UBFC-rPPG session support and stores session-local metadata, RGB samples, frame events, BPM estimates, and analysis events. No new participant data were collected for this report. Any future live recordings require informed consent and appropriate institutional approval.")
add_heading(doc, "Results and Analysis", 3)
add_para(doc, "The available source includes 83 named automated tests across configuration, DSP analysis, timing, ROI extraction, session writing, tracking, calibration, evaluation, UBFC support, and GUI components. The tests document expected behavior for synthetic pulse signals, timing gaps, invalid regions, spectral edge cases, and output schemas. The local project virtual environment is not runnable on this host because its interpreter target is missing; therefore this report does not claim a fresh passing test result or a numerical accuracy result.")
add_heading(doc, "Interpretation and Conclusions on Data", 3)
add_para(doc, "The codebase is structured to support reproducible replay and quality-aware review, but test definitions alone are not execution evidence. A valid final evaluation must run the suite in a repaired environment and report the exact dataset, configuration, reference alignment, accepted-window count, MAE, RMSE, and failure modes. The appropriate conclusion at this stage is readiness for controlled software validation, not verified physiological performance.")
add_heading(doc, "Limitations of Tools (PO(e))", 2)
add_para(doc, "OpenCV timestamp behavior, camera drivers, video codecs, MediaPipe landmark quality, Python dependency versions, and display scheduling can alter observed timing and performance. Classical rPPG is sensitive to lighting and motion. The project handles some failures through validation and logging, but no software-only configuration can remove these dependencies. The current host also demonstrates a deployment limitation: the saved virtual environment refers to an unavailable Python installation.")
add_heading(doc, "Impact Assessment (PO(f))", 2)
add_heading(doc, "Assessment of Societal and Cultural Issues", 3)
add_para(doc, "Contactless sensing may be convenient in educational demonstrations, but facial video is sensitive data. Users should know when processing occurs, what is stored, and how to remove session outputs. Deployment must avoid presenting the system as universally reliable across people or environments without evidence.")
add_heading(doc, "Assessment of Health and Safety Issues", 3)
add_para(doc, "A misleading BPM display can create false reassurance or concern. The software mitigates this by stating its non-medical status, exposing quality information, and rejecting known poor windows. Users must not rely on it in emergencies or clinical decisions.")
add_heading(doc, "Assessment of Legal Issues", 3)
add_para(doc, "Video and physiological inferences may be subject to privacy, consent, institutional, and data-protection requirements. The project should use only authorized camera input and approved datasets, protect session files, and respect dataset licenses and attribution requirements.")
add_heading(doc, "Sustainability Evaluation (PO(g))", 2)
add_para(doc, "A software-only workflow avoids project-specific material waste. Efficient frame handling, bounded buffers, and offline replay reduce unnecessary computation. Sustainability benefits should not be overstated because camera, processor, and display energy use remain part of normal operation.")
add_heading(doc, "Ethical Issues (PO(h))", 2)
add_para(doc, "The design avoids converting a low-confidence observation into a definitive physiological claim. It records rejection causes and labels the output as an estimate. Ethical operation requires informed consent for recorded faces, limited access to session data, honest reporting of validation status, and avoidance of surveillance or health claims outside the demonstrated scope.")

add_heading(doc, "Reflection on Individual and Team Work (PO(i))", 1)
add_heading(doc, "Individual Contribution of Each Member", 2)
add_para(doc, "[Student Name 1] - [Contribution placeholder: DSP design, implementation, testing, documentation].")
add_para(doc, "[Student Name 2] - [Contribution placeholder: GUI, evaluation, testing, documentation].")
add_heading(doc, "Mode of Team Work", 2)
add_para(doc, "The team divided work by software modules, reviewed interfaces and test cases together, and integrated changes through a shared project workspace. Each member should replace this statement with the actual workflow used in the submission.")
add_heading(doc, "Diversity Statement of Team", 2)
add_para(doc, "The team commits to respectful collaboration, equitable task allocation, accessible communication, and evidence-based discussion. Technical decisions are evaluated on documented requirements and testable behavior, not personal characteristics.")
add_heading(doc, "Log Book of Project Implementation", 2)
add_caption(doc, "Table 2: Software project implementation log")
add_table(doc, ["Date", "Milestone achieved", "Individual role", "Team role", "Comments"], [
    ["[Date]", "Requirements and scope defined", "[Name]", "Planning", "Software-only scope confirmed"],
    ["[Date]", "DSP pipeline implemented", "[Name]", "Integration", "Candidate and quality logic added"],
    ["[Date]", "GUI and session logging integrated", "[Name]", "Review", "Replay workflow prepared"],
    ["[Date]", "Tests and report reviewed", "[Name]", "Validation", "Replace with actual activity"],
])

add_heading(doc, "Communication to External Stakeholders (PO(j))", 1)
add_heading(doc, "Executive Summary", 2)
add_para(doc, "This project presents a desktop software tool that estimates a pulse-related frequency from facial video using established signal-processing methods. It tracks facial regions, extracts color traces, evaluates several classical rPPG algorithms, and rejects weak observations instead of always showing a number. The program saves evidence that can be reviewed after a session. It is intended for education and research demonstration only, not for medical use. Before any wider use, the software requires controlled testing with consented data, transparent privacy practices, and validation against an appropriate reference device.")
add_heading(doc, "User Manual", 2)
add_para(doc, "1. Install the project dependencies in a valid Python environment. 2. Launch the desktop application or the command-line runner. 3. Select a permitted live camera, saved video, or supported dataset folder. 4. Keep the face visible with stable lighting and minimal movement. 5. Wait for the warm-up and analysis window to fill. 6. Read BPM only when the status reports an accepted quality result. 7. Stop the session and review the session directory for metadata, CSV records, and evaluation outputs. 8. Do not use the output for diagnosis or urgent decisions.")
add_heading(doc, "Source Repository Link", 2)
add_para(doc, "[Insert authorized source repository URL]")
add_heading(doc, "Demonstration Video Link", 2)
add_para(doc, "[Insert authorized demonstration video URL]")

add_heading(doc, "Project Management (PO(k))", 1)
add_heading(doc, "Timeline of Project Implementation", 2)
add_para(doc, "The project schedule is managed as a software workflow: requirements and literature review; architecture and module design; DSP implementation; UI and session integration; automated tests; replay evaluation; report preparation. The final submission should replace this summary with actual dates and responsibility assignments if required by the course.")
add_caption(doc, "Table 3: Software project timeline")
add_table(doc, ["Phase", "Planned activity", "Deliverable"], [
    ["1", "Requirements and literature review", "Validated scope and references"],
    ["2", "Architecture and DSP module design", "Module interfaces and configuration"],
    ["3", "Implementation and UI integration", "Runnable software workflow"],
    ["4", "Testing and replay evaluation", "Test/evaluation evidence"],
    ["5", "Report review and submission", "Final project report"],
])

add_heading(doc, "Future Work (PO(l))", 1)
add_para(doc, "Future work should first repair and lock the Python environment, execute the automated tests, and publish reproducible dependency instructions. Controlled evaluation can then compare candidates across permitted datasets and documented lighting/motion conditions. Further improvements may include formal latency profiling, a clearer quality dashboard, test coverage for GUI logging parity, configurable privacy retention, and carefully validated learned methods. Any extension toward health use would require substantially stronger evidence, independent validation, privacy review, and appropriate regulatory guidance.")

add_heading(doc, "References", 1)
for ref in [
    "[1] W. Verkruysse, L. O. Svaasand, and J. S. Nelson, \"Remote plethysmographic imaging using ambient light,\" Optics Express, vol. 16, no. 26, pp. 21434-21445, 2008.",
    "[2] G. de Haan and V. Jeanne, \"Robust pulse rate from chrominance-based rPPG,\" IEEE Transactions on Biomedical Engineering, vol. 60, no. 10, pp. 2878-2886, 2013.",
    "[3] W. Wang, A. C. den Brinker, S. Stuijk, and G. de Haan, \"Algorithmic principles of remote PPG,\" IEEE Transactions on Biomedical Engineering, vol. 64, no. 7, pp. 1479-1491, 2017.",
    "[4] M. P. Tarvainen, P. O. Ranta-aho, and P. A. Karjalainen, \"An advanced detrending method with application to HRV analysis,\" IEEE Transactions on Biomedical Engineering, vol. 49, no. 2, pp. 172-175, 2002.",
    "[5] SciPy Community, \"SciPy Signal Processing documentation,\" https://docs.scipy.org/doc/scipy/reference/signal.html, accessed September 2026.",
    "[6] OpenCV, \"OpenCV documentation,\" https://docs.opencv.org/, accessed September 2026.",
]:
    add_para(doc, ref)

# Request Word to refresh fields when the document is opened.
settings = doc.settings.element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")

doc.save(OUTPUT)
print(OUTPUT)
