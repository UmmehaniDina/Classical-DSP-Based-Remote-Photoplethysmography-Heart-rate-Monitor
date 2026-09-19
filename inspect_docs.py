from pathlib import Path

from docx import Document


OUT = Path(r"C:\Users\dumme\Downloads\DSP Project\report_work")
OUT.mkdir(parents=True, exist_ok=True)
for path in [
    Path(r"C:\Users\dumme\Downloads\rPPG_Project_Documentation.docx"),
    Path(r"C:\Users\dumme\Downloads\EEE-xxx-project-report-template (1).docx"),
]:
    doc = Document(path)
    chunks = [f"P {i}: [{para.style.name if para.style else 'None'}] {para.text}" for i, para in enumerate(doc.paragraphs)]
    chunks.append("\nTABLES")
    for index, tbl in enumerate(doc.tables):
        chunks.append(f"TABLE {index}")
        chunks.extend(" | ".join(cell.text for cell in row.cells) for row in tbl.rows)
    out = OUT / f"{path.stem}.txt"
    out.write_text("\n".join(chunks), encoding="utf-8")
    print(path.name, len(doc.paragraphs), len(doc.tables), out)
