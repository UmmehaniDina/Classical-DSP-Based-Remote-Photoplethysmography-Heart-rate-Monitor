from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


REPORT = Path(r"C:\Users\dumme\Downloads\DSP Project\BUET_Classical_rPPG_Final_Project_Report.docx")


def set_text(paragraph, value):
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def insert_after(paragraph, text, style="Normal"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = paragraph._parent.add_paragraph()
    inserted._p.getparent().remove(inserted._p)
    new_p.addprevious(inserted._p)
    inserted.style = style
    inserted.add_run(text)
    return inserted


def header_row(table):
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


doc = Document(REPORT)
p = doc.paragraphs

# 6.1
set_text(p[119], "Individual Contribution of Each Member")
set_text(p[120], "2206080 - Ummehani Dina contributed to the timing and preprocessing strand, including review of capture timestamps, uniform resampling requirements, smoothness-priors detrending, and safe normalization of the RGB traces. She also worked across the application layer by reviewing configuration flow, command-line options, session metadata, and the presentation of rejected conditions. Her work therefore connected the DSP input assumptions with the reproducibility requirements of a usable application.")
set_text(p[121], "2206081 - Md. Montashim Billaha Chisty contributed to landmark-guided ROI handling, motion checks, skin-pixel screening, and the CHROM and POS candidate paths. He also took part in the desktop interaction and telemetry side, particularly the practical flow from frame processing to user-visible status and quality information. This combination helped ensure that the colour-projection algorithms were considered together with the conditions under which a user receives or does not receive an estimate.")
anchor = insert_after(p[121], "2206082 - Mahir Nurain Shawchchow contributed to the spectral-analysis and estimate-decision strand: band selection, Butterworth filtering, Welch PSD inspection, peak-interval agreement, harmonic handling, and quality-aware fusion. He also worked across the runner and evaluation workflow, including session outputs, reference alignment, and the interpretation of saved results. This linked the core BPM decision to an auditable evaluation path rather than treating it as an isolated numerical output.")
anchor = insert_after(anchor, "2206083 - Israt Ferdous Sabrin contributed to candidate-method comparison, including Green and PCA baselines, tracking behaviour, quality-context checks, and regression-oriented review of controlled cases. She also worked across the application and reporting side by considering GUI state, user-facing warm-up or rejection messages, and the consistency of saved evidence. Her contribution helped keep the DSP pipeline interpretable when a signal is weak, incomplete, or unsuitable for a BPM update.")

# 6.2
set_text(p[122], "Mode of Team Work")
set_text(p[123], "The team worked through shared module boundaries rather than treating the project as four isolated parts. At each stage, one member took primary responsibility for examining a DSP or application strand, while the remaining members reviewed its inputs, outputs, configuration assumptions, and interaction with adjacent modules. Development proceeded from input timing and ROI extraction to RGB preprocessing, candidate construction, filtering and estimation, quality control, and application-level logging. Regular short reviews were used to compare the current behaviour against the configured acceptance rules and the expected session records. This arrangement gave every member exposure to both signal-processing decisions and the application, evaluation, and documentation work needed to make those decisions usable.")

# 6.3
set_text(p[124], "Diversity Statement of Team")
set_text(p[125], "The team maintained a respectful working environment in which technical decisions were discussed through reproducible evidence, configuration settings, test cases, and observed outputs. Work was distributed so that no member was confined to a single type of task; each member contributed to a DSP or algorithmic component and to at least one application, evaluation, testing, or documentation activity. Meetings and review notes were kept focused on the task, with space for members to raise concerns about assumptions, unclear behaviour, or incomplete validation. The team recognises that fair participation depends on clear communication, accessible shared materials, and shared responsibility for the final integrated result.")

# 6.4
set_text(p[126], "Log Book of Project Implementation")
table = doc.tables[4]
while len(table.rows) > 1:
    table._tbl.remove(table.rows[-1]._tr)
headers = ["Date", "Milestone and review activity", "Primary contribution", "Team integration and outcome", "Record"]
for cell, value in zip(table.rows[0].cells, headers):
    cell.text = value
header_row(table)

rows = [
    ("20 Jul 2026", "Project proposal accepted; scope refined as a software-only, DSP-centred rPPG monitor.", "All members", "Agreed the data path, non-medical scope, and evidence needed for later evaluation.", "Proposal accepted"),
    ("23 Jul 2026", "Reviewed rPPG literature and converted the problem into requirements for timestamped RGB observations and accepted analysis windows.", "Ummehani Dina; Israt Ferdous Sabrin", "Defined the initial terminology, physiological band, and literature-to-module mapping.", "Requirements note"),
    ("24 Jul 2026", "Mapped the codebase into acquisition, ROI extraction, DSP, tracking, session, evaluation, and desktop-interface responsibilities.", "Md. Montashim Billaha Chisty; Mahir Nurain Shawchchow", "Confirmed interfaces so that DSP results could retain source and quality context.", "Architecture review"),
    ("29 Jul 2026", "Reviewed CaptureClock behaviour, progressing timestamps, duplicate handling, effective FPS, and analysis-window validity checks.", "Ummehani Dina", "Checked how timestamp rules constrain interpolation and downstream filter safety.", "Timing review"),
    ("30 Jul 2026", "Reviewed face-landmark regions, ROI smoothing, normalized motion screening, and fixed YCrCb skin masking.", "Md. Montashim Billaha Chisty", "Linked ROI acceptance and mask coverage to the RGB sample record.", "Extraction review"),
    ("31 Jul 2026", "Reviewed percentile-clipped RGB pooling and the fixed, adaptive-chroma, and temporal-adaptive mask options.", "Israt Ferdous Sabrin; Ummehani Dina", "Checked that invalid samples remain invalid instead of receiving substitute colour values.", "RGB evidence review"),
    ("05 Aug 2026", "Reviewed smoothness-priors detrending, FPS-aware lambda scaling, and channel-wise z-score normalization.", "Ummehani Dina; Mahir Nurain Shawchchow", "Confirmed the preprocessing output required by all candidate methods.", "Preprocessing review"),
    ("06 Aug 2026", "Reviewed Green and PCA baselines and the locally normalised CHROM and POS projections with overlapping windows.", "Israt Ferdous Sabrin; Md. Montashim Billaha Chisty", "Compared candidate metadata and the role of method selection in later quality scoring.", "Candidate-method review"),
    ("13 Aug 2026", "Reviewed Butterworth band-pass processing, Welch PSD settings, peak detection, parabolic refinement, and harmonic adjustment.", "Mahir Nurain Shawchchow", "Connected spectral and time-domain BPM estimates to the configured 55--200 BPM search interval.", "Estimation review"),
    ("20 Aug 2026", "Reviewed multi-feature quality scoring, estimate fusion, jump protection, EMA/Kalman tracking, and rejection categories.", "Mahir Nurain Shawchchow; Israt Ferdous Sabrin", "Checked that low-confidence windows remain visible as rejected or warming up.", "Quality-control review"),
    ("27 Aug 2026", "Reviewed command-line configuration, session folders, CSV schemas, and UBFC-style reference alignment.", "Ummehani Dina; Mahir Nurain Shawchchow", "Prepared the audit trail from accepted samples through BPM estimates to paired evaluation.", "Reproducibility review"),
    ("03 Sep 2026", "Reviewed the background engine, bounded display queues, GUI telemetry, and separation of processing from the Tkinter event loop.", "Md. Montashim Billaha Chisty; Israt Ferdous Sabrin", "Checked the flow from current analysis state to user-visible waveform, spectrum, and status.", "Interactive-flow review"),
    ("10 Sep 2026", "Reviewed controlled tests, saved session outputs, preliminary UBFC-style comparison, and known completion items.", "All members", "Agreed that small-sample results would be reported as preliminary and not as broad accuracy claims.", "Validation review"),
    ("17 Sep 2026", "Completed integration review of the interactive application and prepared the project report evidence.", "All members", "Confirmed the final workflow: input, quality-gated DSP estimate, interface update, and saved session records.", "Interactive app completed"),
]

for values in rows:
    cells = table.add_row().cells
    for cell, value in zip(cells, values):
        cell.text = value

doc.save(REPORT)
print(REPORT)
