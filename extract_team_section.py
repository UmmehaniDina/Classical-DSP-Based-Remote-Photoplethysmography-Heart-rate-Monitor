from copy import deepcopy
from pathlib import Path

from docx import Document


ROOT = Path(r"C:\Users\dumme\Downloads\DSP Project")
SOURCE = ROOT / "BUET_Classical_rPPG_Final_Project_Report.docx"
OUTPUT = ROOT / "Section_6_Reflection_on_Individual_and_Team_Work.docx"

source = Document(SOURCE)
destination = Document()

# Reuse the source style definitions so the extracted section remains consistent
# with the parent report, while keeping the output independent.
for name in ["Normal", "Heading 1", "Heading 2", "Caption"]:
    src_style = source.styles[name]
    dst_style = destination.styles[name]
    dst_style.font.name = src_style.font.name
    dst_style.font.size = src_style.font.size
    dst_style.font.bold = src_style.font.bold
    dst_style.font.italic = src_style.font.italic
    dst_style.paragraph_format.space_before = src_style.paragraph_format.space_before
    dst_style.paragraph_format.space_after = src_style.paragraph_format.space_after
    dst_style.paragraph_format.line_spacing = src_style.paragraph_format.line_spacing

body = destination._element.body
for child in list(body):
    if child.tag.endswith("}sectPr"):
        continue
    body.remove(child)

start = next(i for i, p in enumerate(source.paragraphs) if p.text == "Reflection on Individual and Team Work (PO(i))")
end = next(i for i, p in enumerate(source.paragraphs) if p.text == "Communication to External Stakeholders (PO(j))")

for paragraph in source.paragraphs[start:end]:
    if paragraph.text or paragraph._p.xpath('.//w:drawing'):
        body.insert(len(body) - 1, deepcopy(paragraph._p))

# Copy only the Section 6 logbook table, which is the fifth table in the source.
body.insert(len(body) - 1, deepcopy(source.tables[4]._tbl))

for paragraph in destination.paragraphs:
    if paragraph.text == "Table 4: Software project implementation log":
        if paragraph.runs:
            paragraph.runs[0].text = "Table 1: Software project implementation log"
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run("Table 1: Software project implementation log")

destination.core_properties.title = "Reflection on Individual and Team Work"
destination.core_properties.subject = "Section 6 of the Classical rPPG Monitor final project report"
destination.save(OUTPUT)
print(OUTPUT)
